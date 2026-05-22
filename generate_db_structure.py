# -*- coding: utf-8 -*-
"""
Скрипт для генерации Word-документа со структурой базы данных
Проект: Зори Ваха - Система управления гостиницей
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

print("Начинаю генерацию документа...")

# Создаем документ
doc = Document()

# Заголовок документа
title = doc.add_heading('Структура базы данных', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('Проект: Зори Ваха - Система управления гостиницей')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)

date_p = doc.add_paragraph(f'Дата создания: {datetime.now().strftime("%d.%m.%Y")}')
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.runs[0].font.size = Pt(12)

doc.add_page_break()

print("Создаю таблицу User...")

# Таблица 1: User
doc.add_heading('1. Таблица: User (accounts_user)', level=2)
doc.add_paragraph('Пользователи системы с ролевым доступом')

table = doc.add_table(rows=1, cols=6)
table.style = 'Light Grid Accent 1'

# Заголовки
headers = ['Имя атрибута', 'Описание атрибута', 'Тип', 'Первичный ключ', 'Внешний ключ', 'Ограничения']
hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

# Данные User
user_fields = [
    ['id', 'Уникальный идентификатор', 'UUID', 'Да', '-', 'PRIMARY KEY'],
    ['email', 'Email пользоваостиницы', 'VARCHAR(254)', 'Нет', '-', 'UNIQUE, NOT NULL'],
    ['password', 'Хеш пароля', 'VARCHAR(128)', 'Нет', '-', 'NOT NULL'],
    ['first_name', 'Имя', 'VARCHAR(150)', 'Нет', '-', ''],
    ['last_name', 'Фамилия', 'VARCHAR(150)', 'Нет', '-', ''],
    ['role', 'Роль (user, receptionist, manager, admin, super_admin)', 'VARCHAR(20)', 'Нет', '-', 'DEFAULT user, INDEX'],
    ['phone', 'Телефон', 'VARCHAR(20)', 'Нет', '-', ''],
    ['is_active', 'Активен ли пользоваостиница', 'BOOLEAN', 'Нет', '-', 'DEFAULT TRUE'],
    ['created_at', 'Дата создания', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
    ['updated_at', 'Дата обновления', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
]

for field in user_fields:
    row_cells = table.add_row().cells
    for i, value in enumerate(field):
        row_cells[i].text = value

doc.add_paragraph()

print("Создаю таблицу RoomCategory...")

# Таблица 2: RoomCategory
doc.add_heading('2. Таблица: RoomCategory (hotel_roomcategory)', level=2)
doc.add_paragraph('Категории номеров (Стандарт, Делюкс, Люкс и т.д.)')

table = doc.add_table(rows=1, cols=6)
table.style = 'Light Grid Accent 1'

hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

room_category_fields = [
    ['id', 'ID категории', 'INTEGER', 'Да', '-', 'PRIMARY KEY, AUTO_INCREMENT'],
    ['name', 'Название категории', 'VARCHAR(100)', 'Нет', '-', 'UNIQUE, NOT NULL'],
    ['slug', 'URL slug', 'VARCHAR(120)', 'Нет', '-', 'UNIQUE, NOT NULL'],
    ['description', 'Полное описание', 'TEXT', 'Нет', '-', 'NOT NULL'],
    ['max_guests', 'Максимум гостей', 'SMALLINT', 'Нет', '-', 'DEFAULT 2, >= 1'],
    ['bed_type', 'Тип кровати', 'VARCHAR(20)', 'Нет', '-', ''],
    ['base_price_per_night', 'Базовая цена за ночь', 'DECIMAL(10,2)', 'Нет', '-', 'NOT NULL, >= 0'],
    ['is_active', 'Активна', 'BOOLEAN', 'Нет', '-', 'DEFAULT TRUE, INDEX'],
    ['is_featured', 'Показывать на главной', 'BOOLEAN', 'Нет', '-', 'DEFAULT FALSE, INDEX'],
    ['created_at', 'Дата создания', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
]

for field in room_category_fields:
    row_cells = table.add_row().cells
    for i, value in enumerate(field):
        row_cells[i].text = value

doc.add_paragraph()

print("Создаю таблицу Room...")

# Таблица 3: Room
doc.add_heading('3. Таблица: Room (hotel_room)', level=2)
doc.add_paragraph('Физические номера в гостинице')

table = doc.add_table(rows=1, cols=6)
table.style = 'Light Grid Accent 1'

hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

room_fields = [
    ['id', 'ID номера', 'INTEGER', 'Да', '-', 'PRIMARY KEY, AUTO_INCREMENT'],
    ['category_id', 'Категория номера', 'INTEGER', 'Нет', 'hotel_roomcategory.id', 'NOT NULL, INDEX'],
    ['number', 'Номер комнаты', 'VARCHAR(10)', 'Нет', '-', 'NOT NULL'],
    ['subdivision', 'Подразделение (а, б, в)', 'VARCHAR(5)', 'Нет', '-', ''],
    ['floor', 'Этаж', 'SMALLINT', 'Нет', '-', 'DEFAULT 1'],
    ['max_guests_per_room', 'Макс. персон в номере', 'SMALLINT', 'Нет', '-', 'DEFAULT 1, >= 1'],
    ['status', 'Статус номера', 'VARCHAR(20)', 'Нет', '-', 'INDEX'],
    ['has_balcony', 'Балкон', 'BOOLEAN', 'Нет', '-', 'DEFAULT FALSE'],
    ['has_sea_view', 'Вид на море', 'BOOLEAN', 'Нет', '-', 'DEFAULT FALSE'],
    ['created_at', 'Дата создания', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
]

for field in room_fields:
    row_cells = table.add_row().cells
    for i, value in enumerate(field):
        row_cells[i].text = value

doc.add_paragraph()

print("Создаю таблицу Booking...")

# Таблица 4: Booking
doc.add_heading('4. Таблица: Booking (bookings_booking)', level=2)
doc.add_paragraph('Основная таблица бронирований')

table = doc.add_table(rows=1, cols=6)
table.style = 'Light Grid Accent 1'

hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

booking_fields = [
    ['id', 'ID брони', 'UUID', 'Да', '-', 'PRIMARY KEY'],
    ['confirmation_number', 'Номер брони (ZV-XXXXXXXX)', 'VARCHAR(12)', 'Нет', '-', 'UNIQUE, INDEX, NOT NULL'],
    ['guest_id', 'Гость', 'UUID', 'Нет', 'accounts_user.id', 'NOT NULL'],
    ['room_category_id', 'Категория номера', 'INTEGER', 'Нет', 'hotel_roomcategory.id', 'NOT NULL'],
    ['room_id', 'Назначенный номер', 'INTEGER', 'Нет', 'hotel_room.id', ''],
    ['check_in', 'Дата заезда', 'DATE', 'Нет', '-', 'INDEX, NOT NULL'],
    ['check_out', 'Дата выезда', 'DATE', 'Нет', '-', 'NOT NULL'],
    ['adults', 'Взрослых', 'SMALLINT', 'Нет', '-', 'DEFAULT 1, >= 1'],
    ['children', 'Детей', 'SMALLINT', 'Нет', '-', 'DEFAULT 0'],
    ['total_price', 'Итого', 'DECIMAL(12,2)', 'Нет', '-', 'NOT NULL, >= 0'],
    ['status', 'Статус брони', 'VARCHAR(20)', 'Нет', '-', 'INDEX'],
    ['payment_status', 'Статус оплаты', 'VARCHAR(20)', 'Нет', '-', 'INDEX'],
    ['created_at', 'Дата создания', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
]

for field in booking_fields:
    row_cells = table.add_row().cells
    for i, value in enumerate(field):
        row_cells[i].text = value

doc.add_paragraph()

print("Создаю таблицу Organization...")

# Таблица 5: Organization
doc.add_heading('5. Таблица: Organization (crm_organization)', level=2)
doc.add_paragraph('Организации и корпоративные клиенты')

table = doc.add_table(rows=1, cols=6)
table.style = 'Light Grid Accent 1'

hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

org_fields = [
    ['id', 'ID организации', 'INTEGER', 'Да', '-', 'PRIMARY KEY, AUTO_INCREMENT'],
    ['name', 'Название', 'VARCHAR(255)', 'Нет', '-', 'UNIQUE, NOT NULL'],
    ['org_type', 'Тип организации', 'VARCHAR(20)', 'Нет', '-', 'INDEX'],
    ['inn', 'ИНН', 'VARCHAR(12)', 'Нет', '-', 'INDEX'],
    ['phone', 'Телефон', 'VARCHAR(20)', 'Нет', '-', ''],
    ['email', 'Email', 'VARCHAR(254)', 'Нет', '-', ''],
    ['corporate_discount_pct', 'Корпоративная скидка (%)', 'DECIMAL(5,2)', 'Нет', '-', 'DEFAULT 0'],
    ['is_active', 'Активна', 'BOOLEAN', 'Нет', '-', 'DEFAULT TRUE, INDEX'],
    ['is_approved', 'Подтверждена', 'BOOLEAN', 'Нет', '-', 'DEFAULT FALSE, INDEX'],
    ['created_at', 'Дата создания', 'TIMESTAMP', 'Нет', '-', 'AUTO'],
]

for field in org_fields:
    row_cells = table.add_row().cells
    for i, value in enumerate(field):
        row_cells[i].text = value

# Сохраняем документ
filename = 'Структура_БД_Зори_Ваха.docx'
doc.save(filename)

print(f"✅ Документ успешно создан: {filename}")
print(f"📄 Создано 5 таблиц с описанием структуры БД")
